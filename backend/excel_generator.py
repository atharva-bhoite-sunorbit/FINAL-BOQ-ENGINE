from __future__ import annotations

from io import BytesIO
from typing import Any
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def generate_boq_excel(boq_data: dict[str, Any]) -> bytes:
    wb = openpyxl.Workbook()
    
    # -------------------------------------------------------------
    # Styles & Palettes
    # -------------------------------------------------------------
    title_font = Font(name="Calibri", size=16, bold=True, color="1F4E79")
    subtitle_font = Font(name="Calibri", size=11, italic=True, color="595959")
    section_font = Font(name="Calibri", size=13, bold=True, color="1F4E79")
    header_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=10, color="000000")
    bold_data_font = Font(name="Calibri", size=10, bold=True, color="000000")
    total_font = Font(name="Calibri", size=11, bold=True, color="1F4E79")

    primary_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    teal_fill = PatternFill(start_color="0E6655", end_color="0E6655", fill_type="solid")
    alt_fill = PatternFill(start_color="F2F6FA", end_color="F2F6FA", fill_type="solid")
    kpi_fill = PatternFill(start_color="E8F1F9", end_color="E8F1F9", fill_type="solid")
    total_fill = PatternFill(start_color="D9E2EC", end_color="D9E2EC", fill_type="solid")

    thin_border = Border(
        left=Side(style="thin", color="D3D3D3"),
        right=Side(style="thin", color="D3D3D3"),
        top=Side(style="thin", color="D3D3D3"),
        bottom=Side(style="thin", color="D3D3D3"),
    )
    total_border = Border(
        top=Side(style="thin", color="1F4E79"),
        bottom=Side(style="double", color="1F4E79"),
    )

    currency_fmt = '"₹"#,##0.00'
    number_fmt = '#,##0.00'

    # -------------------------------------------------------------
    # Sheet 1: Project Overview & KPIs
    # -------------------------------------------------------------
    ws_summary = wb.active
    ws_summary.title = "Project Overview"
    ws_summary.views.sheetView[0].showGridLines = True

    ws_summary.append(["KRISALA DEVELOPERS"])
    ws_summary.cell(row=1, column=1).font = title_font
    ws_summary.append(["CONSTRUCTION ESTIMATION & GEOMETRY-FIRST BOQ ENGINE"])
    ws_summary.cell(row=2, column=1).font = subtitle_font
    ws_summary.append([])

    proj = boq_data.get("project", {})
    cost = boq_data.get("cost", {})
    draw = boq_data.get("drawing_summary", {})
    items = boq_data.get("items", [])
    mto_items = boq_data.get("materials", {}).get("items", [])

    ws_summary.append(["Project Drawing File", proj.get("file_name", "N/A")])
    ws_summary.append(["Drawing Units", str(proj.get("units", "mm")).upper()])
    ws_summary.append(["Coordinate Scale", proj.get("scale", "1:1 Standard")])
    ws_summary.append(["Total CAD Entities Parsed", draw.get("entities", len(items))])
    ws_summary.append(["Drawing Layers Detected", len(draw.get("layers", []))])
    ws_summary.append([])

    for r in range(4, 9):
        ws_summary.cell(row=r, column=1).font = bold_data_font
        ws_summary.cell(row=r, column=2).font = data_font

    ws_summary.append(["COST & TAKEOFF EXECUTIVE SUMMARY"])
    ws_summary.cell(row=10, column=1).font = section_font
    ws_summary.append([])

    kpis = [
        ("Total BOQ Line Items", len(items), "nos"),
        ("Total Constituent Materials Required", len(mto_items), "items"),
        ("Material Cost Takeoff", cost.get("material_cost", 0.0), "₹"),
        ("Labour Cost Estimate", cost.get("labour_cost", 0.0), "₹"),
        ("Subtotal (Direct Cost)", cost.get("subtotal", 0.0), "₹"),
        (f"GST / Works Contract Tax ({cost.get('tax_rate', 18)}%)", cost.get("tax", 0.0), "₹"),
        ("Grand Total Project Cost", cost.get("grand_total", 0.0), "₹"),
    ]

    for label, val, unit in kpis:
        r = ws_summary.max_row + 1
        ws_summary.append([label, val])
        cell_lbl = ws_summary.cell(row=r, column=1)
        cell_val = ws_summary.cell(row=r, column=2)
        cell_lbl.fill = kpi_fill
        cell_val.fill = kpi_fill
        cell_lbl.border = thin_border
        cell_val.border = thin_border

        if unit == "₹":
            cell_lbl.font = bold_data_font
            cell_val.font = bold_data_font
            cell_val.number_format = currency_fmt
            cell_val.alignment = Alignment(horizontal="right")
        else:
            cell_lbl.font = data_font
            cell_val.font = bold_data_font
            cell_val.alignment = Alignment(horizontal="right")

    _autofit_columns(ws_summary)

    # -------------------------------------------------------------
    # Sheet 2: Bill of Quantities (BOQ)
    # -------------------------------------------------------------
    ws_boq = wb.create_sheet(title="Bill of Quantities (BOQ)")
    ws_boq.views.sheetView[0].showGridLines = True

    boq_headers = [
        "Sr.", "Item Code", "Category", "Item Description", "Specification / Material",
        "Unit", "Drawing Qty", "Gross Qty", "Opening Deduction", "Net Measured Qty",
        "Wastage %", "Final Qty", "Material Rate", "Labour Rate", "Total Unit Rate",
        "Amount (₹)", "Calculation Formula & Audit Remarks"
    ]

    ws_boq.append(boq_headers)
    for col_num in range(1, len(boq_headers) + 1):
        cell = ws_boq.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = primary_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    start_row = 2
    for idx, item in enumerate(items, start=start_row):
        is_alt = (idx % 2 == 1)
        row_data = [
            item.get("sr_no", idx - 1),
            item.get("code", ""),
            item.get("category", ""),
            item.get("description", ""),
            item.get("material", ""),
            item.get("unit", ""),
            item.get("drawing_quantity", item.get("quantity", 0)),
            item.get("gross_quantity", item.get("quantity", 0)),
            item.get("opening_deduction", 0.0),
            item.get("net_quantity", item.get("quantity", 0)),
            item.get("wastage_pct", 0),
            item.get("final_quantity", 0),
            item.get("material_rate", 0),
            item.get("labour_rate", 0),
            item.get("final_rate", item.get("rate", 0)),
            item.get("amount", 0.0),
            item.get("calculation", item.get("remarks", "")),
        ]
        ws_boq.append(row_data)

        for col_num in range(1, len(row_data) + 1):
            cell = ws_boq.cell(row=idx, column=col_num)
            cell.font = data_font
            cell.border = thin_border
            if is_alt:
                cell.fill = alt_fill

            # Format specific numeric columns
            if col_num in {7, 8, 9, 10, 12}:
                cell.number_format = number_fmt
                cell.alignment = Alignment(horizontal="right")
            elif col_num in {13, 14, 15, 16}:
                cell.number_format = currency_fmt
                cell.alignment = Alignment(horizontal="right")
            elif col_num in {1, 6, 11}:
                cell.alignment = Alignment(horizontal="center")

    # Total Row
    end_row = ws_boq.max_row
    total_row_idx = end_row + 1
    ws_boq.cell(row=total_row_idx, column=4, value="TOTAL DIRECT CONSTRUCTION COST")
    ws_boq.cell(row=total_row_idx, column=4).font = total_font
    ws_boq.cell(row=total_row_idx, column=4).alignment = Alignment(horizontal="right")

    amt_cell = ws_boq.cell(row=total_row_idx, column=16)
    amt_cell.value = f"=SUM(P{start_row}:P{end_row})" if end_row >= start_row else 0
    amt_cell.font = total_font
    amt_cell.number_format = currency_fmt
    amt_cell.alignment = Alignment(horizontal="right")

    for c in range(1, len(boq_headers) + 1):
        cell = ws_boq.cell(row=total_row_idx, column=c)
        cell.border = total_border
        cell.fill = total_fill

    _autofit_columns(ws_boq)

    # -------------------------------------------------------------
    # Sheet 3: Material Takeoff (MTO)
    # -------------------------------------------------------------
    ws_mto = wb.create_sheet(title="Material Takeoff (MTO)")
    ws_mto.views.sheetView[0].showGridLines = True

    mto_headers = [
        "Sr.", "Material Code", "Constituent Material Description", "Category",
        "Associated Work Item", "Base Takeoff Qty", "Base Unit", "Consumption Factor",
        "Wastage %", "Required Material Qty", "Packaging Unit", "Unit Material Rate (₹)",
        "Total Material Cost (₹)"
    ]

    ws_mto.append(mto_headers)
    for col_num in range(1, len(mto_headers) + 1):
        cell = ws_mto.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = teal_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    mto_start_row = 2
    for idx, mat in enumerate(mto_items, start=mto_start_row):
        is_alt = (idx % 2 == 1)
        row_data = [
            mat.get("sr_no", idx - 1),
            mat.get("material_code", ""),
            mat.get("material_name", ""),
            mat.get("category", ""),
            mat.get("derived_from", ""),
            mat.get("base_quantity", 0),
            mat.get("base_unit", ""),
            mat.get("consumption_factor", 1.0),
            mat.get("wastage_percent", 0),
            mat.get("final_quantity", 0),
            mat.get("unit", ""),
            mat.get("unit_rate", 0),
            mat.get("total_cost", 0),
        ]
        ws_mto.append(row_data)

        for col_num in range(1, len(row_data) + 1):
            cell = ws_mto.cell(row=idx, column=col_num)
            cell.font = data_font
            cell.border = thin_border
            if is_alt:
                cell.fill = alt_fill

            if col_num in {6, 8, 10}:
                cell.number_format = number_fmt
                cell.alignment = Alignment(horizontal="right")
            elif col_num in {12, 13}:
                cell.number_format = currency_fmt
                cell.alignment = Alignment(horizontal="right")
            elif col_num in {1, 7, 9, 11}:
                cell.alignment = Alignment(horizontal="center")

    mto_end_row = ws_mto.max_row
    mto_tot_idx = mto_end_row + 1
    ws_mto.cell(row=mto_tot_idx, column=4, value="TOTAL MATERIAL TAKEOFF COST")
    ws_mto.cell(row=mto_tot_idx, column=4).font = total_font
    ws_mto.cell(row=mto_tot_idx, column=4).alignment = Alignment(horizontal="right")

    mto_cost_cell = ws_mto.cell(row=mto_tot_idx, column=13)
    mto_cost_cell.value = f"=SUM(M{mto_start_row}:M{mto_end_row})" if mto_end_row >= mto_start_row else 0
    mto_cost_cell.font = total_font
    mto_cost_cell.number_format = currency_fmt
    mto_cost_cell.alignment = Alignment(horizontal="right")

    for c in range(1, len(mto_headers) + 1):
        cell = ws_mto.cell(row=mto_tot_idx, column=c)
        cell.border = total_border
        cell.fill = total_fill

    _autofit_columns(ws_mto)

    # -------------------------------------------------------------
    # Sheet 4: Geometry Audit Trace
    # -------------------------------------------------------------
    ws_audit = wb.create_sheet(title="Geometry Audit Log")
    ws_audit.views.sheetView[0].showGridLines = True

    audit_headers = [
        "Work Item Code", "Work Category", "Entity Handle", "CAD Entity Type",
        "CAD Layer", "Length / Area", "Height / Thickness (m)", "Assumption Used?"
    ]
    ws_audit.append(audit_headers)
    for col_num in range(1, len(audit_headers) + 1):
        cell = ws_audit.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = primary_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    audit_row = 2
    for item in items:
        code = item.get("code", "")
        cat = item.get("category", "")
        for geom in item.get("source_geometry", []):
            ws_audit.append([
                code,
                cat,
                geom.get("handle") or "N/A",
                geom.get("entity_type", ""),
                geom.get("layer", ""),
                geom.get("area_m2") or geom.get("length_m") or "—",
                geom.get("height_m") or "—",
                "YES" if geom.get("assumption_used") else "NO",
            ])
            for c in range(1, len(audit_headers) + 1):
                cell = ws_audit.cell(row=audit_row, column=c)
                cell.font = data_font
                cell.border = thin_border
            audit_row += 1

    _autofit_columns(ws_audit)

    output = BytesIO()
    wb.save(output)
    return output.getvalue()


def _autofit_columns(ws):
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or "")
            if "\n" in val_str:
                val_str = max(val_str.split("\n"), key=len)
            max_len = max(max_len, len(val_str))
        ws.column_dimensions[col_letter].width = max(10, min(max_len + 3, 50))
