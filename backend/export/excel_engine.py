from __future__ import annotations

from io import BytesIO
from typing import Any
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def generate_boq_excel(boq_data: dict[str, Any]) -> bytes:
    wb = openpyxl.Workbook()

    # Colors & Fonts
    title_font = Font(name="Calibri", size=16, bold=True, color="1F4E79")
    subtitle_font = Font(name="Calibri", size=11, italic=True, color="595959")
    section_font = Font(name="Calibri", size=12, bold=True, color="1F4E79")
    header_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=10, color="000000")
    bold_data_font = Font(name="Calibri", size=10, bold=True, color="000000")
    total_font = Font(name="Calibri", size=11, bold=True, color="1F4E79")

    primary_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    section_fill = PatternFill(start_color="D9E2EC", end_color="D9E2EC", fill_type="solid")
    alt_fill = PatternFill(start_color="F7FAFC", end_color="F7FAFC", fill_type="solid")

    thin_border = Border(
        left=Side(style="thin", color="D3D3D3"),
        right=Side(style="thin", color="D3D3D3"),
        top=Side(style="thin", color="D3D3D3"),
        bottom=Side(style="thin", color="D3D3D3"),
    )
    total_border = Border(
        top=Side(style="thin", color="1F4E79"),
        bottom=Side(style="double", color="1F4E79"),
    )

    currency_fmt = '"₹"#,##0.00'
    number_fmt = '#,##0.00'

    # Extract BOQ Items
    items = boq_data.get("boq", [])
    if not items and "items" in boq_data:
        items = boq_data["items"]

    draw_summary = boq_data.get("drawing_summary", {})

    # -------------------------------------------------------------
    # Sheet 1: Project Overview & Executive Summary
    # -------------------------------------------------------------
    ws_summary = wb.active
    ws_summary.title = "Project Overview"
    ws_summary.views.sheetView[0].showGridLines = True

    ws_summary.append(["CONSTRUCTION ESTIMATION & BOQ ENGINE"])
    ws_summary.cell(row=1, column=1).font = title_font
    ws_summary.append(["CLIENT ESTIMATE & CAD-DRIVEN MATERIAL TAKEOFF"])
    ws_summary.cell(row=2, column=1).font = subtitle_font
    ws_summary.append([])

    ws_summary.append(["Drawing Filename:", draw_summary.get("filename", "Drawing")])
    ws_summary.append(["Drawing Units:", draw_summary.get("units", "m")])
    ws_summary.append(["Total CAD Entities:", draw_summary.get("entity_count", 0)])
    ws_summary.append(["Total Layers Detected:", draw_summary.get("layer_count", 0)])
    ws_summary.append(["Total BOQ Line Items:", len(items)])

    total_amount = sum(float(it.get("amount", 0.0)) for it in items)
    ws_summary.append([])
    ws_summary.append(["Total Estimated Construction Cost:", total_amount])
    ws_summary.cell(row=ws_summary.max_row, column=1).font = bold_data_font
    ws_summary.cell(row=ws_summary.max_row, column=2).font = total_font
    ws_summary.cell(row=ws_summary.max_row, column=2).number_format = currency_fmt

    # Format Overview column widths
    ws_summary.column_dimensions["A"].width = 35
    ws_summary.column_dimensions["B"].width = 30

    # -------------------------------------------------------------
    # Sheet 2: Bill of Quantities (Chronological)
    # -------------------------------------------------------------
    ws_boq = wb.create_sheet(title="Bill of Quantities")
    ws_boq.views.sheetView[0].showGridLines = True

    headers = [
        "Item No", "Section", "Element", "Material", "Description",
        "Unit", "Net Qty", "Wastage %", "Total Qty", "Rate (₹)", "Amount (₹)", "Status"
    ]
    ws_boq.append(headers)
    for col_idx in range(1, len(headers) + 1):
        c = ws_boq.cell(row=1, column=col_idx)
        c.font = header_font
        c.fill = primary_fill
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    current_section = None
    row_num = 2
    for it in items:
        sec = it.get("section", "4. RCC SUPER-STRUCTURE")
        if sec != current_section:
            current_section = sec
            ws_boq.append([current_section] + [""] * (len(headers) - 1))
            sec_cell = ws_boq.cell(row=row_num, column=1)
            sec_cell.font = section_font
            sec_cell.fill = section_fill
            row_num += 1

        net_val = it.get("net_quantity") if it.get("net_quantity") is not None else it.get("quantity")
        tot_val = it.get("total_quantity") if it.get("total_quantity") is not None else it.get("quantity")
        net_q = float(net_val) if net_val is not None else 0.0
        w_pct = float(it.get("wastage_percent", 0.0) or 0.0)
        tot_q = float(tot_val) if tot_val is not None else 0.0
        rate = float(it.get("rate", 0.0) or 0.0)
        amt = float(it.get("amount", tot_q * rate) or 0.0)

        row_vals = [
            it.get("item_no", row_num - 1),
            it.get("section", ""),
            it.get("element_type", ""),
            it.get("material", ""),
            it.get("description", ""),
            it.get("unit", ""),
            net_q,
            w_pct,
            tot_q,
            rate,
            amt,
            it.get("status", "DERIVED"),
        ]
        ws_boq.append(row_vals)

        for c_idx in range(1, len(headers) + 1):
            cell = ws_boq.cell(row=row_num, column=c_idx)
            cell.font = data_font
            cell.border = thin_border
            if row_num % 2 == 1:
                cell.fill = alt_fill

        ws_boq.cell(row=row_num, column=7).number_format = number_fmt
        ws_boq.cell(row=row_num, column=9).number_format = number_fmt
        ws_boq.cell(row=row_num, column=10).number_format = currency_fmt
        ws_boq.cell(row=row_num, column=11).number_format = currency_fmt
        row_num += 1

    # Total row
    ws_boq.append(["", "", "", "", "TOTAL ESTIMATED AMOUNT", "", "", "", "", "", total_amount, ""])
    tot_row = ws_boq.max_row
    for c_idx in range(1, len(headers) + 1):
        cell = ws_boq.cell(row=tot_row, column=c_idx)
        cell.font = total_font
        cell.border = total_border
    ws_boq.cell(row=tot_row, column=11).number_format = currency_fmt

    # Set column widths for BOQ sheet
    widths = [8, 25, 16, 22, 45, 8, 12, 10, 12, 12, 16, 12]
    for i, w in enumerate(widths, start=1):
        ws_boq.column_dimensions[get_column_letter(i)].width = w

    # -------------------------------------------------------------
    # Sheet 3: Material Takeoff (MTO) Summary
    # -------------------------------------------------------------
    ws_mto = wb.create_sheet(title="Material Takeoff (MTO)")
    ws_mto.views.sheetView[0].showGridLines = True

    mto_headers = ["Sr No", "Material Description", "Category", "Unit", "Total Quantity", "Approx Rate (₹)", "Cost (₹)"]
    ws_mto.append(mto_headers)
    for col_idx in range(1, len(mto_headers) + 1):
        c = ws_mto.cell(row=1, column=col_idx)
        c.font = header_font
        c.fill = primary_fill
        c.alignment = Alignment(horizontal="center", vertical="center")

    # Aggregate by material name
    aggregated: dict[str, dict[str, Any]] = {}
    for it in items:
        mat_name = it.get("material", "General")
        u = it.get("unit", "")
        q_val = it.get("total_quantity")
        qty = float(q_val) if q_val is not None else 0.0
        amt = float(it.get("amount", 0.0) or 0.0)
        rate = float(it.get("rate", 0.0) or 0.0)
        if mat_name not in aggregated:
            aggregated[mat_name] = {
                "name": mat_name,
                "category": it.get("section", "General"),
                "unit": u,
                "quantity": 0.0,
                "amount": 0.0,
                "rate": rate,
            }
        aggregated[mat_name]["quantity"] += qty
        aggregated[mat_name]["amount"] += amt

    mto_row = 2
    for idx, (m_key, m_val) in enumerate(aggregated.items(), start=1):
        ws_mto.append([
            idx,
            m_val["name"],
            m_val["category"],
            m_val["unit"],
            round(m_val["quantity"], 3),
            m_val["rate"],
            round(m_val["amount"], 2),
        ])
        for c_idx in range(1, len(mto_headers) + 1):
            cell = ws_mto.cell(row=mto_row, column=c_idx)
            cell.font = data_font
            cell.border = thin_border
        ws_mto.cell(row=mto_row, column=5).number_format = number_fmt
        ws_mto.cell(row=mto_row, column=6).number_format = currency_fmt
        ws_mto.cell(row=mto_row, column=7).number_format = currency_fmt
        mto_row += 1

    mto_widths = [8, 38, 25, 10, 16, 16, 18]
    for i, w in enumerate(mto_widths, start=1):
        ws_mto.column_dimensions[get_column_letter(i)].width = w

    # -------------------------------------------------------------
    # Sheet 4: CAD Traceability & Audit Trail
    # -------------------------------------------------------------
    ws_audit = wb.create_sheet(title="CAD Traceability Audit")
    ws_audit.views.sheetView[0].showGridLines = True

    audit_headers = ["Item No", "Material", "Source Layers", "Source CAD Handles", "Calculation Basis / Formula", "Status", "Confidence"]
    ws_audit.append(audit_headers)
    for col_idx in range(1, len(audit_headers) + 1):
        c = ws_audit.cell(row=1, column=col_idx)
        c.font = header_font
        c.fill = primary_fill

    audit_row = 2
    for it in items:
        handles = it.get("source_entities", [])
        handles_str = ", ".join(handles[:15]) if isinstance(handles, list) else str(handles)
        layers = it.get("source_layers", [])
        layers_str = ", ".join(layers) if isinstance(layers, list) else str(layers)

        ws_audit.append([
            it.get("item_no", audit_row - 1),
            it.get("material", ""),
            layers_str,
            handles_str,
            it.get("calculation_basis", ""),
            it.get("status", "DERIVED"),
            it.get("confidence", 0.90),
        ])
        for c_idx in range(1, len(audit_headers) + 1):
            cell = ws_audit.cell(row=audit_row, column=c_idx)
            cell.font = data_font
            cell.border = thin_border
        audit_row += 1

    audit_widths = [10, 30, 22, 35, 55, 15, 12]
    for i, w in enumerate(audit_widths, start=1):
        ws_audit.column_dimensions[get_column_letter(i)].width = w

    out = BytesIO()
    wb.save(out)
    return out.getvalue()
